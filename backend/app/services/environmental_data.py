from __future__ import annotations

import csv
import hashlib
import json
import math
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, timezone
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any

from uuid import uuid4

from sqlalchemy import select
from sqlalchemy.exc import IntegrityError

from app.config import settings
from app.database.environmental_data import EnvironmentalImport, EnvironmentalObservation
from app.database.urban_memory_store import SessionLocal, initialize_database
from app.schemas.environmental_data import (
    EnvironmentalImportResult,
    METRIC_UNITS,
    EnvironmentalRecord,
    EnvironmentalRecordInput,
    EnvironmentalRecordPage,
    ImportValidationReport,
    UploadPreview,
    ValidationIssue,
)

MAX_PREVIEW_ROWS = 10
REQUIRED_MAPPING_FIELDS = {
    "ward_id", "observation_date", "metric", "value", "unit", "data_classification", "source_name",
}
ALLOWED_MAPPING_FIELDS = REQUIRED_MAPPING_FIELDS | {"source_reference", "notes"}
ENVIRONMENTAL_METRIC_ALIASES = {
    "aqi": "aqi",
    "pm25": "pm25",
    "pm2.5": "pm25",
    "pm_2_5": "pm25",
    "pm10": "pm10",
    "pm10.0": "pm10",
    "no2": "no2",
    "no₂": "no2",
    "o3": "o3",
    "o₃": "o3",
    "co": "co",
    "so2": "so2",
    "so₂": "so2",
    "temperature": "temperature",
    "temperature_c": "temperature",
    "land_surface_temperature": "land_surface_temperature",
    "mean_lst_c": "land_surface_temperature",
    "green_cover": "green_cover",
    "green_cover_percent": "green_cover",
    "population": "population",
    "tot_p": "population",
}


class EnvironmentalUploadError(ValueError):
    pass


class DuplicateEnvironmentalRecord(ValueError):
    pass


@dataclass(frozen=True)
class ParsedEnvironmentalFile:
    filename: str
    checksum: str
    columns: list[str]
    rows: list[tuple[int, dict[str, Any]]]


@dataclass(frozen=True)
class PreparedEnvironmentalImport:
    parsed: ParsedEnvironmentalFile
    mapping: dict[str, str]
    valid_records: list[tuple[int, EnvironmentalRecordInput]]
    errors: list[ValidationIssue]
    duplicate_row_numbers: set[int]

    @property
    def report(self) -> ImportValidationReport:
        duplicates = len(self.duplicate_row_numbers)
        invalid = len({issue.row_number for issue in self.errors if issue.kind == "invalid" and issue.row_number > 1})
        return ImportValidationReport(
            filename=self.parsed.filename,
            checksum=self.parsed.checksum,
            total_rows_read=len(self.parsed.rows),
            valid_rows=len(self.valid_records),
            invalid_rows=invalid,
            duplicate_rows=duplicates,
            errors=self.errors,
        )


def _json_value(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)


def _read_csv(content: bytes) -> tuple[list[str], list[tuple[int, dict[str, Any]]]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise EnvironmentalUploadError("CSV must be UTF-8 encoded.") from exc
    if "\x00" in text:
        raise EnvironmentalUploadError("CSV contains unsupported binary data.")
    reader = csv.reader(StringIO(text), strict=True)
    try:
        header_row = next(reader)
    except (StopIteration, csv.Error) as exc:
        raise EnvironmentalUploadError("CSV is empty or malformed.") from exc
    headers = [cell.strip() for cell in header_row]
    if not headers or any(not header for header in headers):
        raise EnvironmentalUploadError("CSV must have non-empty column headers.")
    if len(headers) > 100 or len({header.casefold() for header in headers}) != len(headers):
        raise EnvironmentalUploadError("CSV headers must be unique and contain no more than 100 columns.")

    rows: list[tuple[int, dict[str, Any]]] = []
    try:
        for row_number, values in enumerate(reader, start=2):
            if not any(str(value).strip() for value in values):
                continue
            if len(rows) >= settings.ENVIRONMENTAL_UPLOAD_MAX_ROWS:
                raise EnvironmentalUploadError(
                    f"File exceeds the {settings.ENVIRONMENTAL_UPLOAD_MAX_ROWS}-row limit."
                )
            if len(values) != len(headers):
                raise EnvironmentalUploadError(f"CSV row {row_number} has a different number of columns than the header.")
            rows.append((row_number, dict(zip(headers, values))))
    except csv.Error as exc:
        raise EnvironmentalUploadError("CSV is malformed.") from exc
    return headers, rows


def _read_xlsx(content: bytes) -> tuple[list[str], list[tuple[int, dict[str, Any]]]]:
    if not content.startswith(b"PK"):
        raise EnvironmentalUploadError("File content is not a valid XLSX workbook.")
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > settings.ENVIRONMENTAL_UPLOAD_MAX_BYTES * 20:
                raise EnvironmentalUploadError("XLSX expands beyond the supported uncompressed size limit.")
    except zipfile.BadZipFile as exc:
        raise EnvironmentalUploadError("File content is not a valid XLSX workbook.") from exc
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
    except Exception as exc:
        raise EnvironmentalUploadError("XLSX workbook is malformed or unsupported.") from exc
    try:
        worksheet = workbook.active
        iterator = worksheet.iter_rows()
        try:
            header_cells = next(iterator)
        except StopIteration as exc:
            raise EnvironmentalUploadError("XLSX worksheet is empty.") from exc
        if any(cell.data_type == "f" for cell in header_cells):
            raise EnvironmentalUploadError("XLSX headers cannot contain formulas.")
        headers = [str(cell.value).strip() if cell.value is not None else "" for cell in header_cells]
        while headers and not headers[-1]:
            headers.pop()
        if not headers or any(not header for header in headers):
            raise EnvironmentalUploadError("XLSX must have non-empty column headers.")
        if len(headers) > 100 or len({header.casefold() for header in headers}) != len(headers):
            raise EnvironmentalUploadError("XLSX headers must be unique and contain no more than 100 columns.")

        rows: list[tuple[int, dict[str, Any]]] = []
        for row_number, cells in enumerate(iterator, start=2):
            if any(cell.value is not None and str(cell.value).strip() for cell in cells[len(headers):]):
                raise EnvironmentalUploadError(f"XLSX row {row_number} contains data beyond the header columns.")
            cells = cells[:len(headers)]
            if any(cell.data_type == "f" for cell in cells):
                raise EnvironmentalUploadError(f"XLSX row {row_number} contains a formula; formulas are not evaluated or imported.")
            values = [cell.value for cell in cells]
            if not any(value is not None and str(value).strip() for value in values):
                continue
            if len(rows) >= settings.ENVIRONMENTAL_UPLOAD_MAX_ROWS:
                raise EnvironmentalUploadError(
                    f"File exceeds the {settings.ENVIRONMENTAL_UPLOAD_MAX_ROWS}-row limit."
                )
            if len(values) < len(headers):
                values.extend([None] * (len(headers) - len(values)))
            rows.append((row_number, dict(zip(headers, values))))
        return headers, rows
    finally:
        workbook.close()


def parse_environmental_file(content: bytes, filename: str) -> ParsedEnvironmentalFile:
    if not filename or len(filename) > 255:
        raise EnvironmentalUploadError("A valid filename is required.")
    if len(content) > settings.ENVIRONMENTAL_UPLOAD_MAX_BYTES:
        raise EnvironmentalUploadError(
            f"File exceeds the {settings.ENVIRONMENTAL_UPLOAD_MAX_BYTES}-byte upload limit."
        )
    suffix = Path(filename).suffix.casefold()
    if suffix == ".csv":
        columns, rows = _read_csv(content)
    elif suffix == ".xlsx":
        columns, rows = _read_xlsx(content)
    elif suffix == ".xls":
        raise EnvironmentalUploadError("Legacy .xls files are not supported; export as .xlsx or .csv.")
    else:
        raise EnvironmentalUploadError("Only .csv and .xlsx files are supported.")
    return ParsedEnvironmentalFile(
        filename=Path(filename).name,
        checksum=hashlib.sha256(content).hexdigest(),
        columns=columns,
        rows=rows,
    )


def get_known_ward_ids() -> set[int]:
    snapshot = Path(__file__).resolve().parents[1] / "digital_twin" / "delhi_virtual_city_model.json"
    try:
        with snapshot.open("r", encoding="utf-8") as handle:
            model = json.load(handle)
    except (OSError, json.JSONDecodeError):
        return set()
    return {
        int(ward["ward_no"])
        for ward in model.get("wards", [])
        if isinstance(ward, dict) and str(ward.get("ward_no", "")).isdigit()
    }


def _normalize_metric(value: Any) -> str:
    key = str(value or "").strip().casefold().replace(" ", "_")
    metric = ENVIRONMENTAL_METRIC_ALIASES.get(key, key)
    if metric not in METRIC_UNITS:
        raise ValueError(f"Unsupported metric '{value}'.")
    return metric


def _normalize_unit(metric: str, value: Any) -> str:
    raw = str(value or "").strip().replace("μ", "µ")
    key = raw.casefold().replace(" ", "").replace("³", "3")
    aliases = {
        "aqi": "AQI",
        "µg/m3": "µg/m³",
        "ug/m3": "µg/m³",
        "mg/m3": "mg/m³",
        "degc": "°C",
        "c": "°C",
        "percent": "%",
        "person": "people",
        "persons": "people",
    }
    unit = aliases.get(key, raw)
    if unit not in METRIC_UNITS[metric]:
        raise ValueError(f"Unit '{raw}' is not supported for metric '{metric}'.")
    return unit


def _parse_date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    if not text:
        raise ValueError("Observation date is required.")
    try:
        return date.fromisoformat(text[:10])
    except ValueError as exc:
        raise ValueError("Observation date must be an ISO date or timestamp.") from exc


def _row_record(row: dict[str, Any], mapping: dict[str, str], known_wards: set[int]) -> EnvironmentalRecordInput:
    ward_raw = row.get(mapping["ward_id"])
    try:
        ward_number = float(ward_raw)
    except (TypeError, ValueError) as exc:
        raise ValueError("Ward ID must be a known whole-number ward identifier.") from exc
    if not math.isfinite(ward_number) or not ward_number.is_integer():
        raise ValueError("Ward ID must be a known whole-number ward identifier.")
    ward_id = int(ward_number)
    if ward_id not in known_wards:
        raise ValueError(f"Unknown ward ID '{ward_raw}'.")

    metric = _normalize_metric(row.get(mapping["metric"]))
    unit = _normalize_unit(metric, row.get(mapping["unit"]))
    try:
        value = float(row.get(mapping["value"]))
    except (TypeError, ValueError) as exc:
        raise ValueError("Value must be numeric.") from exc
    mapped = {
        "ward_id": ward_id,
        "observation_date": _parse_date(row.get(mapping["observation_date"])),
        "metric": metric,
        "value": value,
        "unit": unit,
        "data_classification": str(row.get(mapping["data_classification"]) or "").strip().casefold(),
        "source_name": str(row.get(mapping["source_name"]) or "").strip(),
        "source_reference": str(row.get(mapping["source_reference"]) or "").strip() or None
        if "source_reference" in mapping else None,
        "notes": str(row.get(mapping["notes"]) or "").strip() or None if "notes" in mapping else None,
    }
    return EnvironmentalRecordInput.model_validate(mapped)


def _record_key(record: EnvironmentalRecordInput) -> tuple[int, date, str, str, str]:
    return (
        record.ward_id,
        record.observation_date,
        record.metric,
        record.source_name,
        record.source_reference or "",
    )


def _existing_record_keys(records: list[EnvironmentalRecordInput]) -> set[tuple[int, date, str, str, str]]:
    if not records:
        return set()
    initialize_database()
    wards = {record.ward_id for record in records}
    metrics = {record.metric for record in records}
    dates = [record.observation_date for record in records]
    with SessionLocal() as session:
        rows = session.execute(
            select(
                EnvironmentalObservation.ward_id,
                EnvironmentalObservation.observation_date,
                EnvironmentalObservation.metric,
                EnvironmentalObservation.source_name,
                EnvironmentalObservation.source_reference,
            ).where(
                EnvironmentalObservation.ward_id.in_(wards),
                EnvironmentalObservation.metric.in_(metrics),
                EnvironmentalObservation.observation_date >= min(dates),
                EnvironmentalObservation.observation_date <= max(dates),
            )
        ).all()
    return {
        (row.ward_id, row.observation_date, row.metric, row.source_name, row.source_reference)
        for row in rows
    }


def prepare_environmental_import(
    content: bytes,
    filename: str,
    mapping: dict[str, str],
) -> PreparedEnvironmentalImport:
    parsed = parse_environmental_file(content, filename)
    unknown_fields = set(mapping) - ALLOWED_MAPPING_FIELDS
    missing_fields = REQUIRED_MAPPING_FIELDS - set(mapping)
    invalid_columns = {column for column in mapping.values() if column not in parsed.columns}
    errors: list[ValidationIssue] = []
    mapping_problems = [
        *(f"Unsupported canonical field '{field}'." for field in sorted(unknown_fields)),
        *(f"Required canonical field '{field}' is not mapped." for field in sorted(missing_fields)),
        *(f"Source column '{column}' does not exist." for column in sorted(invalid_columns)),
    ]
    if len(set(mapping.values())) != len(mapping.values()):
        mapping_problems.append("Each canonical field must map to a distinct source column.")
    if mapping_problems:
        row_numbers = [row_number for row_number, _ in parsed.rows] or [1]
        for row_number in row_numbers:
            errors.extend(
                ValidationIssue(row_number=row_number, field="mapping", reason=problem)
                for problem in mapping_problems
            )
    if not parsed.rows:
        errors.append(ValidationIssue(row_number=1, field="rows", reason="File contains no data rows."))
    if errors:
        return PreparedEnvironmentalImport(parsed, mapping, [], errors, set())

    known_wards = get_known_ward_ids()
    valid_records: list[tuple[int, EnvironmentalRecordInput]] = []
    seen: set[tuple[int, date, str, str, str]] = set()
    duplicate_row_numbers: set[int] = set()
    for row_number, raw_row in parsed.rows:
        try:
            record = _row_record(raw_row, mapping, known_wards)
        except Exception as exc:
            errors.append(ValidationIssue(row_number=row_number, field="record", reason=str(exc)))
            continue
        key = _record_key(record)
        valid_records.append((row_number, record))
        if key in seen:
            duplicate_row_numbers.add(row_number)
            errors.append(ValidationIssue(row_number=row_number, field="record", reason="Duplicate record within this file.", kind="duplicate"))
        else:
            seen.add(key)

    existing_keys = _existing_record_keys([record for _, record in valid_records])
    for row_number, record in valid_records:
        if _record_key(record) in existing_keys:
            duplicate_row_numbers.add(row_number)
            errors.append(ValidationIssue(row_number=row_number, field="record", reason="Record already exists for this ward, date, metric, and source.", kind="duplicate"))
    return PreparedEnvironmentalImport(parsed, mapping, valid_records, errors, duplicate_row_numbers)


def preview_environmental_file(content: bytes, filename: str) -> UploadPreview:
    parsed = parse_environmental_file(content, filename)
    return UploadPreview(
        filename=parsed.filename,
        checksum=parsed.checksum,
        columns=parsed.columns,
        preview_rows=[
            {column: _json_value(value) for column, value in row.items()}
            for _, row in parsed.rows[:MAX_PREVIEW_ROWS]
        ],
        total_rows_read=len(parsed.rows),
    )


def validate_environmental_file(content: bytes, filename: str, mapping: dict[str, str]) -> ImportValidationReport:
    return prepare_environmental_import(content, filename, mapping).report


def import_environmental_file(
    content: bytes,
    filename: str,
    mapping: dict[str, str],
    expected_checksum: str,
) -> EnvironmentalImportResult:
    prepared = prepare_environmental_import(content, filename, mapping)
    parsed = prepared.parsed
    if parsed.checksum != expected_checksum:
        raise EnvironmentalUploadError("The selected file changed after preview; preview it again before importing.")
    mapping_digest = hashlib.sha256(json.dumps(mapping, sort_keys=True, separators=(",", ":")).encode()).hexdigest()
    import_id = hashlib.sha256(f"{parsed.checksum}:{mapping_digest}".encode()).hexdigest()
    initialize_database()
    with SessionLocal() as session:
        existing_import = session.get(EnvironmentalImport, import_id)
        if existing_import is not None:
            return EnvironmentalImportResult(
                filename=existing_import.filename,
                checksum=existing_import.checksum,
                total_rows_read=existing_import.total_rows_read,
                valid_rows=existing_import.valid_rows,
                invalid_rows=existing_import.invalid_rows,
                duplicate_rows=existing_import.duplicate_rows,
                errors=prepared.errors,
                import_id=existing_import.import_id,
                imported_at=existing_import.created_at,
                rows_inserted=0,
                rows_skipped=existing_import.total_rows_read,
                idempotent_retry=True,
            )

        insertion_candidates = [
            (row_number, record)
            for row_number, record in prepared.valid_records
            if row_number not in prepared.duplicate_row_numbers
        ]
        try:
            session.add(EnvironmentalImport(
                import_id=import_id,
                checksum=parsed.checksum,
                mapping_digest=mapping_digest,
                filename=parsed.filename,
                total_rows_read=prepared.report.total_rows_read,
                valid_rows=prepared.report.valid_rows,
                invalid_rows=prepared.report.invalid_rows,
                duplicate_rows=prepared.report.duplicate_rows,
                rows_inserted=len(insertion_candidates),
                rows_skipped=prepared.report.invalid_rows + prepared.report.duplicate_rows,
            ))
            for row_number, record in insertion_candidates:
                session.add(EnvironmentalObservation(
                    record_id=f"ENV-{hashlib.sha256(f'{import_id}:{row_number}'.encode()).hexdigest()[:16].upper()}",
                    ward_id=record.ward_id,
                    observation_date=record.observation_date,
                    metric=record.metric,
                    value=record.value,
                    unit=record.unit,
                    data_classification=record.data_classification,
                    source_name=record.source_name,
                    source_reference=record.source_reference or "",
                    notes=record.notes,
                    source_import_id=import_id,
                    source_row_number=row_number,
                ))
            session.commit()
        except IntegrityError as exc:
            session.rollback()
            raise DuplicateEnvironmentalRecord("Import conflicts with records saved concurrently; validate again.") from exc

    return EnvironmentalImportResult(
        **prepared.report.model_dump(),
        import_id=import_id,
        imported_at=datetime.now(timezone.utc),
        rows_inserted=len(insertion_candidates),
        rows_skipped=prepared.report.invalid_rows + prepared.report.duplicate_rows,
    )


def create_environmental_record(record: EnvironmentalRecordInput) -> EnvironmentalRecord:
    if record.ward_id not in get_known_ward_ids():
        raise ValueError(f"Unknown ward ID '{record.ward_id}'.")
    initialize_database()
    with SessionLocal() as session:
        observation = EnvironmentalObservation(
            record_id=f"ENV-{__import__('uuid').uuid4().hex[:16].upper()}",
            ward_id=record.ward_id,
            observation_date=record.observation_date,
            metric=record.metric,
            value=record.value,
            unit=record.unit,
            data_classification=record.data_classification,
            source_name=record.source_name,
            source_reference=record.source_reference or "",
            notes=record.notes,
        )
        session.add(observation)
        try:
            session.commit()
        except IntegrityError as exc:
            session.rollback()
            raise DuplicateEnvironmentalRecord("A record with this ward, date, metric, and source already exists.") from exc
        session.refresh(observation)
        return _to_record_schema(observation)


def _to_record_schema(record: EnvironmentalObservation) -> EnvironmentalRecord:
    return EnvironmentalRecord(
        record_id=record.record_id,
        ward_id=record.ward_id,
        observation_date=record.observation_date,
        metric=record.metric,
        value=record.value,
        unit=record.unit,
        data_classification=record.data_classification,
        source_name=record.source_name,
        source_reference=record.source_reference or None,
        notes=record.notes,
        created_at=record.created_at,
    )


def list_environmental_records(
    *, limit: int, offset: int, ward_id: int | None = None, metric: str | None = None,
) -> EnvironmentalRecordPage:
    initialize_database()
    with SessionLocal() as session:
        query = session.query(EnvironmentalObservation)
        if ward_id is not None:
            query = query.filter(EnvironmentalObservation.ward_id == ward_id)
        if metric is not None:
            query = query.filter(EnvironmentalObservation.metric == metric)
        total = query.count()
        records = query.order_by(
            EnvironmentalObservation.observation_date.desc(),
            EnvironmentalObservation.created_at.desc(),
            EnvironmentalObservation.record_id,
        ).offset(offset).limit(limit).all()
        return EnvironmentalRecordPage(
            records=[_to_record_schema(record) for record in records],
            total=total,
            limit=limit,
            offset=offset,
        )